"""
export_to_excel.py
==================
Exports irs990_2023_2024.db to Excel workbook for boss presentation.

Sheets:
  1. Summary          — key stats at a glance
  2. Jewish Grants    — all 90,560 Jewish grants (full data)
  3. By Year          — grant counts and dollars by year
  4. By Subject       — top subject areas
  5. Top Grantmakers  — top 50 grantmakers by total dollars
  6. Top Recipients   — top 50 recipients by total dollars

Note: 3.7M rows cannot fit in one Excel sheet (1M row limit).
      Full data exported as CSV instead. Excel has the key subsets.
"""

import sqlite3
import sys
from pathlib import Path

import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from project_utils.paths import BOSS_DELIVERABLES_2023_2024_DIR, WAREHOUSE_DB_PATH

BASE_DIR = Path(__file__).parent
DB_PATH = WAREHOUSE_DB_PATH
OUT_DIR = BOSS_DELIVERABLES_2023_2024_DIR
OUT_DIR.mkdir(parents=True, exist_ok=True)

EXCEL_OUT = OUT_DIR / "IRS_Form990_Grants_2023_2024.xlsx"
CSV_FULL  = OUT_DIR / "IRS_Grants_2023_2024_FULL.csv"
CSV_JEWS  = OUT_DIR / "Jewish_Grants_2023_2024.csv"

print("Connecting to database...")
conn = sqlite3.connect(str(DB_PATH))

# ── Load data ──────────────────────────────────────────────────────────────────
print("Loading Jewish grants...")
df_jews = pd.read_sql("""
    SELECT
        tax_year        AS "Tax Year",
        grantmaker_name AS "Grantmaker Name",
        grantmaker_state AS "Grantmaker State",
        recipient_name  AS "Recipient Name",
        recipient_city  AS "Recipient City",
        recipient_state AS "Recipient State",
        geography       AS "Geography",
        grant_amount    AS "Grant Amount ($)",
        primary_subject AS "Primary Subject",
        support_strategy AS "Support Strategy",
        description     AS "Description",
        ein             AS "EIN"
    FROM irs_grants_2023_2024
    WHERE expense_type = 'Jews'
    ORDER BY tax_year, grant_amount DESC
""", conn)

print(f"  Jewish grants: {len(df_jews):,}")

print("Loading summary stats...")
df_by_year = pd.read_sql("""
    SELECT
        tax_year            AS "Tax Year",
        expense_type        AS "Category",
        COUNT(*)            AS "Grant Count",
        SUM(grant_amount)   AS "Total Amount ($)",
        AVG(grant_amount)   AS "Average Grant ($)",
        MAX(grant_amount)   AS "Largest Grant ($)"
    FROM irs_grants_2023_2024
    GROUP BY tax_year, expense_type
    ORDER BY tax_year, expense_type
""", conn)

df_by_subject = pd.read_sql("""
    SELECT
        primary_subject     AS "Subject Area",
        COUNT(*)            AS "Grant Count",
        SUM(grant_amount)   AS "Total Amount ($)",
        AVG(grant_amount)   AS "Average Grant ($)"
    FROM irs_grants_2023_2024
    WHERE primary_subject IS NOT NULL
    GROUP BY primary_subject
    ORDER BY COUNT(*) DESC
""", conn)

df_top_grantmakers = pd.read_sql("""
    SELECT
        grantmaker_name     AS "Grantmaker Name",
        grantmaker_state    AS "State",
        COUNT(*)            AS "Grant Count",
        SUM(grant_amount)   AS "Total Disbursed ($)",
        AVG(grant_amount)   AS "Average Grant ($)"
    FROM irs_grants_2023_2024
    WHERE grantmaker_name IS NOT NULL
    GROUP BY grantmaker_name
    ORDER BY SUM(grant_amount) DESC
    LIMIT 50
""", conn)

df_top_recipients = pd.read_sql("""
    SELECT
        recipient_name      AS "Recipient Name",
        recipient_state     AS "State",
        COUNT(*)            AS "Times Received",
        SUM(grant_amount)   AS "Total Received ($)",
        AVG(grant_amount)   AS "Average Grant ($)"
    FROM irs_grants_2023_2024
    WHERE recipient_name IS NOT NULL
    GROUP BY recipient_name
    ORDER BY SUM(grant_amount) DESC
    LIMIT 50
""", conn)

df_summary_stats = pd.read_sql("""
    SELECT
        'Total Grants (2023-2024)'  AS "Metric",
        COUNT(*)                    AS "Value"
    FROM irs_grants_2023_2024
    UNION ALL
    SELECT 'Jewish Grants', COUNT(*) FROM irs_grants_2023_2024 WHERE expense_type='Jews'
    UNION ALL
    SELECT 'Non-Jewish Grants', COUNT(*) FROM irs_grants_2023_2024 WHERE expense_type='Non-Jews'
    UNION ALL
    SELECT 'Total Dollars ($)', SUM(grant_amount) FROM irs_grants_2023_2024
    UNION ALL
    SELECT 'Jewish Dollars ($)', SUM(grant_amount) FROM irs_grants_2023_2024 WHERE expense_type='Jews'
    UNION ALL
    SELECT 'Non-Jewish Dollars ($)', SUM(grant_amount) FROM irs_grants_2023_2024 WHERE expense_type='Non-Jews'
    UNION ALL
    SELECT 'Average Grant ($)', AVG(grant_amount) FROM irs_grants_2023_2024
    UNION ALL
    SELECT 'Largest Single Grant ($)', MAX(grant_amount) FROM irs_grants_2023_2024
    UNION ALL
    SELECT 'Unique Grantmakers', COUNT(DISTINCT grantmaker_name) FROM irs_grants_2023_2024
    UNION ALL
    SELECT 'Unique Recipients', COUNT(DISTINCT recipient_name) FROM irs_grants_2023_2024
""", conn)

conn.close()

# ── Format numbers ─────────────────────────────────────────────────────────────
def fmt_dollars(df, cols):
    for c in cols:
        if c in df.columns:
            df[c] = df[c].apply(lambda x: round(x) if pd.notna(x) else None)
    return df

df_jews           = fmt_dollars(df_jews,           ["Grant Amount ($)"])
df_by_year        = fmt_dollars(df_by_year,        ["Total Amount ($)", "Average Grant ($)", "Largest Grant ($)"])
df_by_subject     = fmt_dollars(df_by_subject,     ["Total Amount ($)", "Average Grant ($)"])
df_top_grantmakers= fmt_dollars(df_top_grantmakers,["Total Disbursed ($)", "Average Grant ($)"])
df_top_recipients = fmt_dollars(df_top_recipients, ["Total Received ($)", "Average Grant ($)"])

# ── Write Excel ────────────────────────────────────────────────────────────────
print(f"Writing Excel workbook: {EXCEL_OUT.name} ...")

HEADER_FILL  = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT  = Font(bold=True, color="FFFFFF", size=11)
JEWISH_FILL  = PatternFill("solid", fgColor="FFF2CC")
ALT_FILL     = PatternFill("solid", fgColor="EBF3FB")
BORDER       = Border(
    bottom=Side(style="thin", color="CCCCCC"),
    right=Side(style="thin", color="CCCCCC"),
)

def style_sheet(ws, df, header_fill=HEADER_FILL):
    # Header row
    for cell in ws[1]:
        cell.font      = HEADER_FONT
        cell.fill      = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border    = BORDER

    # Data rows — alternate shading
    for i, row in enumerate(ws.iter_rows(min_row=2), start=2):
        fill = ALT_FILL if i % 2 == 0 else PatternFill()
        for cell in row:
            cell.fill      = fill
            cell.alignment = Alignment(vertical="center", wrap_text=False)
            cell.border    = BORDER

    # Auto-width columns
    for col_idx, col in enumerate(ws.columns, start=1):
        max_len = max((len(str(c.value or "")) for c in col), default=10)
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 4, 45)

    # Freeze header row
    ws.freeze_panes = "A2"


with pd.ExcelWriter(str(EXCEL_OUT), engine="openpyxl") as writer:

    # Sheet 1: Summary
    df_summary_stats.to_excel(writer, sheet_name="Summary", index=False)
    ws = writer.sheets["Summary"]
    ws.sheet_properties.tabColor = "1F4E79"
    style_sheet(ws, df_summary_stats)
    ws["A1"].value = "Metric"
    ws["B1"].value = "Value"
    print("  Sheet 1/6: Summary done")

    # Sheet 2: Jewish Grants
    df_jews.to_excel(writer, sheet_name="Jewish Grants", index=False)
    ws = writer.sheets["Jewish Grants"]
    ws.sheet_properties.tabColor = "FFD700"
    style_sheet(ws, df_jews, PatternFill("solid", fgColor="1F4E79"))
    # Highlight Jewish rows in yellow
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
            cell.fill = JEWISH_FILL
    print(f"  Sheet 2/6: Jewish Grants ({len(df_jews):,} rows) done")

    # Sheet 3: By Year & Category
    df_by_year.to_excel(writer, sheet_name="By Year & Category", index=False)
    ws = writer.sheets["By Year & Category"]
    ws.sheet_properties.tabColor = "70AD47"
    style_sheet(ws, df_by_year)
    print("  Sheet 3/6: By Year & Category done")

    # Sheet 4: By Subject
    df_by_subject.to_excel(writer, sheet_name="By Subject Area", index=False)
    ws = writer.sheets["By Subject Area"]
    ws.sheet_properties.tabColor = "ED7D31"
    style_sheet(ws, df_by_subject)
    print("  Sheet 4/6: By Subject Area done")

    # Sheet 5: Top Grantmakers
    df_top_grantmakers.to_excel(writer, sheet_name="Top 50 Grantmakers", index=False)
    ws = writer.sheets["Top 50 Grantmakers"]
    ws.sheet_properties.tabColor = "5B9BD5"
    style_sheet(ws, df_top_grantmakers)
    print("  Sheet 5/6: Top 50 Grantmakers done")

    # Sheet 6: Top Recipients
    df_top_recipients.to_excel(writer, sheet_name="Top 50 Recipients", index=False)
    ws = writer.sheets["Top 50 Recipients"]
    ws.sheet_properties.tabColor = "A9D18E"
    style_sheet(ws, df_top_recipients)
    print("  Sheet 6/6: Top 50 Recipients done")

print(f"\nExcel saved: {EXCEL_OUT}")
print(f"Excel size : {EXCEL_OUT.stat().st_size/1e6:.1f} MB")

# ── Export CSVs ────────────────────────────────────────────────────────────────
print("\nExporting Jewish grants CSV...")
df_jews.to_csv(str(CSV_JEWS), index=False)
print(f"  Saved: {CSV_JEWS.name} ({CSV_JEWS.stat().st_size/1e6:.1f} MB)")

print("\nExporting full dataset CSV (all 3.7M rows)...")
conn2 = sqlite3.connect(str(DB_PATH))
chunk_iter = pd.read_sql("""
    SELECT
        tax_year, expense_type, grantmaker_name, grantmaker_state,
        recipient_name, recipient_city, recipient_state, geography,
        grant_amount, primary_subject, support_strategy, description, ein
    FROM irs_grants_2023_2024
    ORDER BY tax_year, expense_type
""", conn2, chunksize=50000)

first = True
for chunk in chunk_iter:
    chunk.to_csv(str(CSV_FULL), mode='w' if first else 'a',
                 header=first, index=False)
    first = False

conn2.close()
print(f"  Saved: {CSV_FULL.name} ({CSV_FULL.stat().st_size/1e6:.0f} MB)")

print("\n" + "="*60)
print("ALL EXPORTS COMPLETE")
print("="*60)
print(f"Output folder: {OUT_DIR}")
print(f"  Excel workbook : {EXCEL_OUT.name}")
print(f"  Jewish CSV     : {CSV_JEWS.name}")
print(f"  Full CSV       : {CSV_FULL.name}")
print("="*60)
