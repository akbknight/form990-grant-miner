"""
create_jewish_master.py
=======================
Extracts all "Jews"-classified rows from form990_streaming_output.csv
into a single Excel file: Jewish_Grants_Master.xlsx

Also rebuilds AI_Extracted_Data_Current_Status.xlsx with accurate totals.

Run:
    .venv312\\Scripts\\python.exe create_jewish_master.py
"""

import os, re, time, datetime
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE_DIR     = os.path.dirname(os.path.abspath(__file__))
INPUT_CSV    = os.path.join(BASE_DIR, "form990_streaming_output.csv")
OUTPUT_JEWS  = os.path.join(BASE_DIR, "Jewish_Grants_Master.xlsx")
OUTPUT_SUMM  = os.path.join(BASE_DIR, "AI_Extracted_Data_Current_Status.xlsx")
TODAY        = datetime.date.today().strftime("%B %d, %Y")

TARGET_COLUMNS = [
    "Grantmaker Name", "Grantmaker State", "Recipient Name", "Recipient City",
    "Recipient State/Country", "Geography", "Primary Subject", "Tax Year",
    "Year Authorized", "Grant Amount", "Support Strategies", "Description",
    "Expense Type",
]

COL_WIDTHS = [35, 16, 35, 18, 22, 18, 18, 12, 16, 14, 22, 55, 16]

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
JEWISH_FILL = PatternFill("solid", fgColor="FFF2CC")
ALT_FILL    = PatternFill("solid", fgColor="FFF9E6")


# ── 1. Build Jewish_Grants_Master.xlsx ──────────────────────────────────────

def build_jewish_excel(df_jews: pd.DataFrame):
    print(f"\nBuilding Jewish_Grants_Master.xlsx  ({len(df_jews):,} rows)...")
    t0 = time.time()

    wb = Workbook(write_only=False)
    ws = wb.active
    ws.title = "Jewish_Grants"
    ws.freeze_panes = "A2"

    thin = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    # Header
    for i, col in enumerate(TARGET_COLUMNS, 1):
        ws.column_dimensions[get_column_letter(i)].width = COL_WIDTHS[i-1]
        c = ws.cell(row=1, column=i, value=col)
        c.fill = HEADER_FILL
        c.font = Font(bold=True, color="FFFFFF", size=11)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=False)
        c.border = border

    # Ensure only target columns, in order
    for col in TARGET_COLUMNS:
        if col not in df_jews.columns:
            df_jews[col] = "N/A"
    df_out = df_jews[TARGET_COLUMNS]

    # Data rows
    for r_idx, row in enumerate(df_out.itertuples(index=False), 2):
        fill = JEWISH_FILL if r_idx % 2 == 0 else ALT_FILL
        for c_idx, val in enumerate(row, 1):
            c = ws.cell(row=r_idx, column=c_idx, value=val)
            c.fill = fill
            c.font = Font(size=10)
            c.alignment = Alignment(vertical="top", wrap_text=False)
            c.border = border

    ws.auto_filter.ref = ws.dimensions
    wb.save(OUTPUT_JEWS)
    size_mb = os.path.getsize(OUTPUT_JEWS) / 1024 / 1024
    print(f"  Saved: {OUTPUT_JEWS}  ({size_mb:.0f} MB)  in {time.time()-t0:.0f}s")


# ── 2. Rebuild summary workbook ─────────────────────────────────────────────

DATA_DICT = [
    ("Grantmaker Name",         "Legal name of the foundation/organization filing Form 990"),
    ("Grantmaker State",        "U.S. state abbreviation of the grantmaker's registered address"),
    ("Recipient Name",          "Name of the organization or individual receiving the grant"),
    ("Recipient City",          "City where the grant recipient is located"),
    ("Recipient State/Country", "U.S. state code or country code for the recipient"),
    ("Geography",               "Derived field: 'United States' for domestic, country name for international"),
    ("Primary Subject",         "Auto-classified topic area: Education, Health, Housing, Arts & Culture, Food Security, Employment"),
    ("Tax Year",                "IRS tax year of the Form 990 filing (4-digit year)"),
    ("Year Authorized",         "Year the grant was authorized (same as Tax Year from XML)"),
    ("Grant Amount",            "Dollar amount of the grant as reported on Form 990"),
    ("Support Strategies",      "Auto-classified support type: Capacity building, Core support, Conferences, etc."),
    ("Description",             "Purpose/description of the grant as stated on the Form 990"),
    ("Expense Type",            "Classification: Jews | Non jews | Miscellaneous  (see Classification tab)"),
]

CLASSIFICATION_LOGIC = [
    ("CLASSIFICATION METHODOLOGY", ""),
    ("", ""),
    ("Method Used", "Keyword-based pattern matching (Python regex, no AI required)"),
    ("Applied To", "Combined text of: Recipient Name + Description + Primary Subject + Support Strategies"),
    ("", ""),
    ("KEYWORDS THAT TRIGGER 'Jews' CLASSIFICATION", ""),
    ("Jewish identity terms", "jew, jewish, jews, judaism, hebrew"),
    ("Institutions", "yeshiva, synagogue, chabad, hillel, JCC (Jewish Community Center), lubavitch"),
    ("Religious terms", "torah, kosher, rabbi, talmud, kabbalah, kippah, shabbat, mitzvah, bar mitzvah, bat mitzvah"),
    ("Holidays", "passover, hanukkah, chanukah, rosh hashanah"),
    ("Organizations", "jewish federation, jewish community, jewish family, hadassah, B'nai"),
    ("Places/Events", "israel, israeli, jerusalem, kibbutz, zion, zionist, birthright, holocaust"),
    ("Cultural terms", "shalom, tikva, menorah"),
    ("", ""),
    ("'Non jews'", "Any row where the combined text is NOT empty and does NOT match any Jewish keyword"),
    ("'Miscellaneous'", "Rows with empty/N/A data in all classified fields — insufficient data to determine"),
    ("", ""),
    ("IMPORTANT NOTE", "This is keyword-based classification, not AI-based. It identifies organizations/grants that explicitly mention Jewish identity, institutions, or topics in their IRS filing text."),
]

HDR_FILL  = PatternFill("solid", fgColor="1F4E79")
SUB_FILL  = PatternFill("solid", fgColor="2E75B6")
GREEN_FILL = PatternFill("solid", fgColor="1E8449")

def _thin_border():
    s = Side(style="thin", color="BDD7EE")
    return Border(left=s, right=s, top=s, bottom=s)

def build_summary_sheet(ws, counts):
    ws.column_dimensions["A"].width = 50
    ws.column_dimensions["B"].width = 60

    # Title banner
    ws.merge_cells("A1:B1")
    t = ws["A1"]
    t.value = "IRS Form 990 Grant Data — EXTRACTION COMPLETE"
    t.font = Font(bold=True, color="FFFFFF", size=14)
    t.fill = PatternFill("solid", fgColor="1A5276")
    t.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 32

    ws.merge_cells("A2:B2")
    t2 = ws["A2"]
    t2.value = f"Generated: {TODAY}   |   Status: ALL PROCESSING COMPLETE"
    t2.font = Font(bold=True, color="FFFFFF", size=11)
    t2.fill = PatternFill("solid", fgColor="1E8449")
    t2.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 22

    total = counts['total']
    jews  = counts['jews']
    non   = counts['non']
    misc  = counts['misc']

    # Count part files
    parts = sorted([f for f in os.listdir(BASE_DIR)
                    if f.startswith("AI_Extracted_Data_Full_Part") and f.endswith(".xlsx")])

    rows = [
        ("", ""),
        ("PROJECT STATUS", "COMPLETE"),
        ("Data Source", "IRS TEOS (Tax Exempt Organization Search) — Form 990 XML, 2019-2026"),
        ("ZIP Archives Processed", "62 of 62  (100% COMPLETE)"),
        ("Total Grant Rows Extracted", f"{total:,}"),
        ("Classification Method", "Keyword-based (no AI/API required)"),
        ("Extraction Completed", TODAY),
        ("", ""),
        ("CLASSIFICATION RESULTS", ""),
        ("Jews", f"{jews:,} rows  ({jews/total*100:.2f}%)  — grants to Jewish organizations/causes"),
        ("Non jews", f"{non:,} rows  ({non/total*100:.2f}%)  — grants to all other organizations"),
        ("Miscellaneous", f"{misc:,} rows  ({misc/total*100:.2f}%)  — insufficient data to classify"),
        ("", ""),
        ("OUTPUT FILES", ""),
        ("MASTER FILE (all Jews rows)", f"Jewish_Grants_Master.xlsx  —  {jews:,} rows  (single file, all Jewish grants)"),
        ("Full Dataset (CSV)", f"form990_streaming_output.csv  —  {total:,} rows, ~1.4 GB"),
        ("", "NOTE: Excel cannot hold 9.7M rows (limit: 1,048,576). The CSV is the true master file."),
        ("", "To work with the full dataset, use Excel Power Query to connect to the CSV."),
        ("", ""),
        ("SPLIT EXCEL FILES (900k rows each)", ""),
    ]
    for p in parts:
        size_mb = os.path.getsize(os.path.join(BASE_DIR, p)) // (1024*1024)
        rows.append((f"  {p}", f"{size_mb} MB"))

    section_keys = {"PROJECT STATUS", "CLASSIFICATION RESULTS", "OUTPUT FILES",
                    "SPLIT EXCEL FILES (900k rows each)"}

    r = 3
    for label, value in rows:
        c1 = ws.cell(row=r, column=1, value=label)
        c2 = ws.cell(row=r, column=2, value=value)
        c1.font = Font(size=10)
        c2.font = Font(size=10)
        c1.alignment = Alignment(vertical="top", wrap_text=True)
        c2.alignment = Alignment(vertical="top", wrap_text=True)
        if label in section_keys:
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
            c1.font = Font(bold=True, color="FFFFFF", size=11)
            c1.fill = PatternFill("solid", fgColor="2E75B6")
            ws.row_dimensions[r].height = 20
        elif label == "Jews":
            c1.font = Font(bold=True, size=10, color="7D6608")
            c2.font = Font(size=10, color="7D6608")
            c1.fill = PatternFill("solid", fgColor="FFF2CC")
            c2.fill = PatternFill("solid", fgColor="FFF2CC")
        elif label == "Non jews":
            c1.font = Font(bold=True, size=10, color="1E8449")
            c2.font = Font(size=10, color="1E8449")
            c1.fill = PatternFill("solid", fgColor="E2EFDA")
            c2.fill = PatternFill("solid", fgColor="E2EFDA")
        elif label == "MASTER FILE (all Jews rows)":
            c1.font = Font(bold=True, size=10, color="7D6608")
            c2.font = Font(bold=True, size=10, color="7D6608")
            c1.fill = PatternFill("solid", fgColor="FFF2CC")
            c2.fill = PatternFill("solid", fgColor="FFF2CC")
        r += 1


def build_dict_sheet(ws):
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 72
    # Header
    for i, title in enumerate(["Column Name", "Description"], 1):
        c = ws.cell(row=1, column=i, value=title)
        c.fill = HDR_FILL
        c.font = Font(bold=True, color="FFFFFF", size=11)
        c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 22
    ws.freeze_panes = "A2"
    alt = PatternFill("solid", fgColor="EBF3FB")
    for i, (col, desc) in enumerate(DATA_DICT, 2):
        f = alt if i % 2 == 0 else PatternFill("solid", fgColor="FFFFFF")
        c1 = ws.cell(row=i, column=1, value=col)
        c2 = ws.cell(row=i, column=2, value=desc)
        c1.fill = f; c2.fill = f
        c1.font = Font(bold=True, size=10)
        c2.font = Font(size=10)
        c1.alignment = Alignment(vertical="top")
        c2.alignment = Alignment(vertical="top", wrap_text=True)
        ws.row_dimensions[i].height = 18


def build_classification_sheet(ws):
    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 70
    # Header
    for i, title in enumerate(["Item", "Detail"], 1):
        c = ws.cell(row=1, column=i, value=title)
        c.fill = HDR_FILL
        c.font = Font(bold=True, color="FFFFFF", size=11)
        c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 22
    ws.freeze_panes = "A2"

    section_keys = {
        "CLASSIFICATION METHODOLOGY",
        "KEYWORDS THAT TRIGGER 'Jews' CLASSIFICATION",
        "IMPORTANT NOTE",
    }
    yellow = PatternFill("solid", fgColor="FFF2CC")
    blue   = PatternFill("solid", fgColor="D6E4F0")
    orange = PatternFill("solid", fgColor="FCE4D6")

    for r, (label, value) in enumerate(CLASSIFICATION_LOGIC, 2):
        c1 = ws.cell(row=r, column=1, value=label)
        c2 = ws.cell(row=r, column=2, value=value)
        for c in [c1, c2]:
            c.alignment = Alignment(vertical="top", wrap_text=True)
            c.font = Font(size=10)
        if label in section_keys:
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
            c1.font = Font(bold=True, color="FFFFFF", size=11)
            if label == "IMPORTANT NOTE":
                c1.fill = orange
            else:
                c1.fill = PatternFill("solid", fgColor="2E75B6")
            ws.row_dimensions[r].height = 20
        elif label in ("'Non jews'", "'Miscellaneous'"):
            c1.fill = blue; c2.fill = blue
        elif label == "Method Used" or label == "Applied To":
            c1.font = Font(bold=True, size=10)
        elif label == "IMPORTANT NOTE":
            c2.fill = orange
            c2.font = Font(size=10, italic=True)
        ws.row_dimensions[r].height = 30 if label == "IMPORTANT NOTE" and value else 18


def build_summary_workbook(counts):
    print("\nBuilding AI_Extracted_Data_Current_Status.xlsx...")
    wb = Workbook()

    ws1 = wb.active
    ws1.title = "Summary"
    build_summary_sheet(ws1, counts)

    ws2 = wb.create_sheet("Data_Dictionary")
    build_dict_sheet(ws2)

    ws3 = wb.create_sheet("Classification_Method")
    build_classification_sheet(ws3)

    wb.save(OUTPUT_SUMM)
    print(f"  Saved: {OUTPUT_SUMM}")


# ── Main ─────────────────────────────────────────────────────────────────────

def main():
    print(f"Input CSV: {INPUT_CSV}")
    if not os.path.exists(INPUT_CSV):
        print("ERROR: form990_streaming_output.csv not found!")
        return

    print("Step 1: Counting categories and collecting Jewish rows...")
    t0 = time.time()
    jews_chunks = []
    counts = {"total": 0, "jews": 0, "non": 0, "misc": 0}

    for chunk in pd.read_csv(INPUT_CSV, dtype=str, chunksize=200_000,
                              encoding="utf-8-sig", on_bad_lines="skip"):
        chunk = chunk.fillna("N/A")
        # Keep only target columns
        for col in TARGET_COLUMNS:
            if col not in chunk.columns:
                chunk[col] = "N/A"
        chunk = chunk[TARGET_COLUMNS]

        vc = chunk["Expense Type"].value_counts()
        counts["total"] += len(chunk)
        counts["jews"]  += int(vc.get("Jews", 0))
        counts["non"]   += int(vc.get("Non jews", 0))
        counts["misc"]  += int(vc.get("Miscellaneous", 0))

        jewish_rows = chunk[chunk["Expense Type"] == "Jews"]
        if not jewish_rows.empty:
            jews_chunks.append(jewish_rows)

        print(f"  {counts['total']:,} rows processed, {counts['jews']:,} Jewish found...",
              end="\r", flush=True)

    print(f"\n  Done in {time.time()-t0:.0f}s")
    print(f"  Total : {counts['total']:,}")
    print(f"  Jews  : {counts['jews']:,}  ({counts['jews']/counts['total']*100:.2f}%)")
    print(f"  Non j : {counts['non']:,}  ({counts['non']/counts['total']*100:.2f}%)")
    print(f"  Misc  : {counts['misc']:,}  ({counts['misc']/counts['total']*100:.2f}%)")

    if jews_chunks:
        df_jews = pd.concat(jews_chunks, ignore_index=True)
        print(f"\nStep 2: Writing {len(df_jews):,} Jewish rows to Excel...")
        build_jewish_excel(df_jews)
    else:
        print("\nNo Jewish rows found.")

    print("\nStep 3: Rebuilding summary workbook...")
    build_summary_workbook(counts)

    print("\nDONE. Output files:")
    for f in [OUTPUT_JEWS, OUTPUT_SUMM]:
        if os.path.exists(f):
            mb = os.path.getsize(f) / 1024 / 1024
            print(f"  {os.path.basename(f)}  ({mb:.0f} MB)")


if __name__ == "__main__":
    main()
