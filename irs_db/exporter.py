"""
irs_db/exporter.py
==================
Secure export utilities for the PostgreSQL warehouse.
"""

from __future__ import annotations

import argparse
import logging
import sys
import time
from pathlib import Path
from typing import Optional

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from irs_db import config, loader
from project_utils.export_filters import build_safe_where_clause, ensure_allowed_table
from project_utils.paths import EXPORT_DIR, JEWISH_MASTER_DB_XLSX_PATH

log = logging.getLogger(__name__)
logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s  %(levelname)-8s  %(message)s",
    handlers=[logging.StreamHandler(sys.stdout)],
)

OUT_DIR = EXPORT_DIR
OUT_DIR.mkdir(parents=True, exist_ok=True)
HDR_FILL = PatternFill("solid", fgColor="1F4E79")
JEWISH_FILL = PatternFill("solid", fgColor="FFF2CC")
ALT_FILL = PatternFill("solid", fgColor="FFF9E6")
EXCEL_COLUMNS = [
    "ein",
    "tax_year",
    "grantmaker_name",
    "grantmaker_state",
    "recipient_name",
    "recipient_city",
    "recipient_state_country",
    "geography",
    "primary_subject",
    "support_strategy",
    "grant_amount",
    "description",
    "expense_type",
]
EXCEL_WIDTHS = [11, 10, 35, 14, 35, 18, 22, 18, 20, 28, 14, 55, 16]


def export_csv(
    output_path: str,
    filter_expr: str = "",
    table_or_view: str = "irs_grants",
) -> None:
    """
    Export a safe slice of a supported table/view using PostgreSQL COPY.

    Filters are parsed through an allowlist instead of interpolating raw SQL.
    """
    target = ensure_allowed_table(table_or_view)
    where_sql, params = build_safe_where_clause(filter_expr)

    conn = loader.get_conn()
    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    started_at = time.time()

    base_query = f"SELECT * FROM {target}"
    if where_sql:
        base_query += f" WHERE {where_sql}"
    if target in {"irs_grants", "v_grants_powerbi", "v_filing_summary"}:
        base_query += " ORDER BY id"

    with conn.cursor() as cur, open(output, "w", encoding="utf-8-sig", newline="") as handle:
        copy_sql = cur.mogrify(
            f"COPY ({base_query}) TO STDOUT WITH CSV HEADER",
            params,
        ).decode("utf-8")
        cur.copy_expert(copy_sql, handle)

    size_mb = output.stat().st_size / 1024 / 1024
    log.info("Saved %s (%.0f MB, %.0fs)", output, size_mb, time.time() - started_at)
    conn.close()


def export_all_csv() -> None:
    export_csv(str(OUT_DIR / "export_all_grants.csv"), table_or_view="v_grants_powerbi")


def export_jewish_csv() -> None:
    export_csv(str(OUT_DIR / "export_jewish_grants.csv"), filter_expr="expense_type='Jews'")


def export_summary_csv() -> None:
    export_csv(str(OUT_DIR / "export_filing_summary.csv"), table_or_view="v_filing_summary")


def export_filtered_csv(filter_expr: str) -> None:
    safe_name = filter_expr.replace("/", "_").replace(" ", "_").replace("'", "")[:40]
    export_csv(str(OUT_DIR / f"export_filtered_{safe_name}.csv"), filter_expr=filter_expr)


def export_jewish_excel(output_path: Optional[str] = None) -> None:
    """Export Jewish grants to a styled Excel workbook."""
    output = Path(output_path or JEWISH_MASTER_DB_XLSX_PATH)
    conn = loader.get_conn()

    started_at = time.time()
    query = f"""
        SELECT {', '.join(EXCEL_COLUMNS)}
        FROM irs_grants
        WHERE expense_type = 'Jews'
        ORDER BY tax_year, grantmaker_name, grant_amount DESC
    """
    frame = pd.read_sql(query, conn)
    conn.close()

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Jewish_Grants"
    sheet.freeze_panes = "A2"

    thin = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for index, column in enumerate(EXCEL_COLUMNS, 1):
        sheet.column_dimensions[get_column_letter(index)].width = EXCEL_WIDTHS[index - 1]
        cell = sheet.cell(row=1, column=index, value=column.replace("_", " ").title())
        cell.fill = HDR_FILL
        cell.font = Font(bold=True, color="FFFFFF", size=11)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border
    sheet.row_dimensions[1].height = 22

    for row_index, row in enumerate(frame.itertuples(index=False), 2):
        fill = JEWISH_FILL if row_index % 2 == 0 else ALT_FILL
        for column_index, value in enumerate(row, 1):
            cell = sheet.cell(row=row_index, column=column_index, value=value)
            cell.fill = fill
            cell.font = Font(size=10)
            cell.alignment = Alignment(vertical="top")
            cell.border = border

    sheet.auto_filter.ref = sheet.dimensions
    workbook.save(str(output))
    log.info("Saved %s in %.0fs", output, time.time() - started_at)


def create_powerbi_views() -> None:
    """Ensure the reporting views exist in the warehouse."""
    conn = loader.get_conn()
    views_sql = """
    CREATE OR REPLACE VIEW v_grants_powerbi AS
    SELECT
        g.id, g.ein,
        COALESCE(f.org_name, g.grantmaker_name) AS grantmaker_name,
        g.grantmaker_state, g.tax_year,
        g.recipient_name, g.recipient_city, g.recipient_state_country,
        g.geography, g.primary_subject, g.support_strategy,
        g.grant_amount, g.expense_type, g.description,
        f.form_type, f.total_revenue, f.total_expenses,
        f.total_assets_eoy, f.net_assets_eoy, f.employee_count,
        f.org_city, f.org_zip
    FROM irs_grants g
    LEFT JOIN irs_filings f ON f.ein = g.ein AND f.tax_year = g.tax_year;

    CREATE OR REPLACE VIEW v_filing_summary AS
    SELECT
        f.id, f.ein, f.tax_year, f.form_type, f.org_name, f.org_state, f.org_city,
        f.total_revenue, f.total_expenses, f.total_assets_eoy, f.net_assets_eoy,
        f.employee_count,
        COUNT(g.id) AS total_grants,
        SUM(g.grant_amount) AS total_grant_dollars,
        COUNT(g.id) FILTER (WHERE g.expense_type = 'Jews') AS jewish_grant_count,
        SUM(g.grant_amount) FILTER (WHERE g.expense_type = 'Jews') AS jewish_grant_dollars,
        COUNT(g.id) FILTER (WHERE g.expense_type = 'Non jews') AS non_jewish_grant_count,
        SUM(g.grant_amount) FILTER (WHERE g.expense_type = 'Non jews') AS non_jewish_grant_dollars
    FROM irs_filings f
    LEFT JOIN irs_grants g ON g.ein = f.ein AND g.tax_year = f.tax_year
    GROUP BY f.id;

    CREATE OR REPLACE VIEW v_stats AS
    SELECT
        tax_year,
        expense_type,
        COUNT(*) AS grant_count,
        SUM(grant_amount) AS total_dollars,
        AVG(grant_amount) AS avg_grant,
        MAX(grant_amount) AS max_grant
    FROM irs_grants
    GROUP BY tax_year, expense_type
    ORDER BY tax_year, expense_type;
    """
    with conn.cursor() as cur:
        cur.execute(views_sql)
    conn.commit()
    conn.close()
    log.info("Power BI views created: v_grants_powerbi, v_filing_summary, v_stats")


def main() -> None:
    parser = argparse.ArgumentParser(description="Export IRS data from PostgreSQL")
    parser.add_argument("--jewish-excel", action="store_true", help="Export Jewish grants to Excel")
    parser.add_argument("--jewish-csv", action="store_true", help="Export Jewish grants to CSV")
    parser.add_argument("--all", action="store_true", help="Export all grants to CSV")
    parser.add_argument("--summary-csv", action="store_true", help="Export filing summary to CSV")
    parser.add_argument("--csv", action="store_true", help="Export a filtered CSV (use with --filter)")
    parser.add_argument(
        "--filter",
        type=str,
        default="",
        help=(
            "Safe filter syntax such as tax_year=2023, expense_type='Jews', "
            "grantmaker_state=NY, or grantmaker_name~foundation"
        ),
    )
    parser.add_argument("--powerbi-views", action="store_true", help="Create Power BI views in the DB")
    parser.add_argument("--output", type=str, help="Custom output file path")
    args = parser.parse_args()

    try:
        if args.jewish_excel:
            export_jewish_excel(args.output)
        elif args.jewish_csv:
            export_jewish_csv()
        elif args.all:
            export_all_csv()
        elif args.summary_csv:
            export_summary_csv()
        elif args.csv:
            if not args.filter:
                raise ValueError("--csv requires --filter. Example: --filter \"expense_type='Jews'\"")
            export_filtered_csv(args.filter)
        elif args.powerbi_views:
            create_powerbi_views()
        else:
            parser.print_help()
    except ValueError as exc:
        log.error(str(exc))
        sys.exit(1)


if __name__ == "__main__":
    main()
